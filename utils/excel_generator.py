import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from config import Config

def generate_custom_excel(session_id, documents_data, selected_attribute_keys, output_filename=None):
    """
    Generate a professionally formatted, editable Excel (.xlsx) file containing ONLY selected attributes.
    
    Args:
        session_id (int): ID of current processing session
        documents_data (list of dict): List of document dicts with extracted fields
        selected_attribute_keys (list of str): List of attribute keys chosen by the user
        output_filename (str, optional): Custom filename
        
    Returns:
        tuple: (full_file_path, base_filename, total_records, total_columns)
    """
    os.makedirs(Config.GENERATED_EXCEL_FOLDER, exist_ok=True)
    
    if not output_filename:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_filename = f"extracted_documents_session_{session_id}_{timestamp}.xlsx"
        
    file_path = os.path.join(Config.GENERATED_EXCEL_FOLDER, output_filename)
    
    # Map attribute keys to clean labels
    attr_labels = {attr['id']: attr['label'] for attr in Config.AVAILABLE_ATTRIBUTES}
    
    # If no attributes selected, default to common ones
    if not selected_attribute_keys:
        selected_attribute_keys = ['name', 'dob', 'income', 'community', 'certificate_no']
        
    headers = [attr_labels.get(k, k.replace('_', ' ').title()) for k in selected_attribute_keys]
    
    # Build rows of data
    rows = []
    for doc in documents_data:
        row = {}
        fields = doc.get('fields', {})
        for k in selected_attribute_keys:
            val = fields.get(k, '')
            if isinstance(val, dict):
                val = val.get('value', '')
            row[attr_labels.get(k, k.replace('_', ' ').title())] = val
        rows.append(row)
        
    # Create OpenPyXL workbook for high-quality custom styling
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extracted Data"
    
    # Style definitions
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep navy
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    data_font = Font(name="Calibri", size=11, color="1F2937")
    data_alignment = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    # 1. Write Header Row
    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
    
    ws.row_dimensions[1].height = 28
    
    # 2. Write Data Rows
    for row_idx, row_data in enumerate(rows, start=2):
        ws.row_dimensions[row_idx].height = 22
        is_even = (row_idx % 2 == 0)
        
        for col_idx, key in enumerate(headers, start=1):
            val = row_data.get(key, '')
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.alignment = data_alignment
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill
                
    # 3. Freeze Header Row (so header stays visible when scrolling)
    ws.freeze_panes = 'A2'
    
    # 4. Enable Auto-filter on all columns
    if headers:
        max_col_letter = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{max_col_letter}{max(2, len(rows) + 1)}"
    
    # 5. Auto-fit column widths with safety margins
    for col_idx, header_text in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(header_text))
        
        for row_idx in range(2, len(rows) + 2):
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val is not None:
                max_len = max(max_len, len(str(cell_val)))
                
        # Apply comfortable width with bounds (min 15, max 45)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 15), 45)
        
    wb.save(file_path)
    wb.close()
    
    return file_path, output_filename, len(rows), len(headers)
